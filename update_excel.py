import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def update_excel(file_path):
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)
    
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Dark blue
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=11, bold=True)
    regular_font = Font(name="Calibri", size=11)
    money_format = "$#,##0"
    
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    
    # ----------------------------------------------------
    # HOJA 1: Compras
    # ----------------------------------------------------
    ws1 = wb.create_sheet(title="Compras")
    
    # Titulos
    ws1["A1"] = "Total primera compra (Ignacio)"
    ws1["B1"] = 122251
    ws1["A2"] = "Total segunda compra (Ignacio)"
    ws1["B2"] = 61141
    ws1["A3"] = "Total tercera compra (Mindy)"
    ws1["B3"] = 11790
    ws1["A4"] = "Total cuarta compra (Joaquín - Pan choripán)"
    ws1["B4"] = 5570
    ws1["A5"] = "Total quinta compra (Joaquín - Empanadas)"
    ws1["B5"] = 46250
    
    for r in range(1, 6):
        ws1[f"A{r}"].font = bold_font
        ws1[f"B{r}"].font = bold_font
        ws1[f"B{r}"].number_format = money_format
    
    items = [
        "Chorizo", "Lomo Liso", "Lomo vetado", "Costillar Chileno", "Costillar cerdo",
        "Trutro", "Pechuga", "Carbón", "Pan Choripán", "Trigo Mote (1kg)", "Huesillos (1kg)", "Chancaca y Azúcar", "Restante",
        "Pino (1a Empanada)", "Pino (2a Empanada)", "Mechada", "Camarón Queso", "Queso", "Napolitana"
    ]
    
    c1_values = [14670, 21226, 19560, 19238, 14957, 0, 0, 0, 0, 2490, 10990, 0, 19120, 0, 0, 0, 0, 0, 0]
    c2_values = [4890, 0, 33693, 0, 0, 16353, 6205, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0]
    c3_values = [0, 0, 0, 0, 0, 0, 0, 8400, 0, 0, 0, 3390, 0, 0, 0, 0, 0, 0, 0]
    c4_values = [0, 0, 0, 0, 0, 0, 0, 0, 5570, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0]
    c5_values = [0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 17500, 15000, 6000, 3000, 2250, 2500]
    
    # Headers
    header_row_idx = 7
    ws1.cell(row=header_row_idx, column=1, value="Detalle").font = header_font
    ws1.cell(row=header_row_idx, column=1).fill = header_fill
    ws1.cell(row=header_row_idx, column=1).alignment = Alignment(horizontal="center")
    
    for col_idx, item in enumerate(items, start=2):
        cell = ws1.cell(row=header_row_idx, column=col_idx, value=item)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    
    total_col_idx = len(items) + 2
    cell_tot = ws1.cell(row=header_row_idx, column=total_col_idx, value="Total Fila")
    cell_tot.font = header_font
    cell_tot.fill = header_fill
    cell_tot.alignment = Alignment(horizontal="center")
    
    # Rows
    rows_data = [
        ("Primera compra (Ignacio)", c1_values, "=B1"),
        ("Segunda compra (Ignacio)", c2_values, "=B2"),
        ("Tercera compra (Mindy)", c3_values, "=B3"),
        ("Cuarta compra (Joaquín)", c4_values, "=B4"),
        ("Quinta compra (Joaquín - Empanadas)", c5_values, "=B5"),
    ]
    
    for r_offset, (label, vals, expected_ref) in enumerate(rows_data, start=8):
        ws1.cell(row=r_offset, column=1, value=label).font = bold_font
        for col_idx, val in enumerate(vals, start=2):
            c = ws1.cell(row=r_offset, column=col_idx, value=val)
            c.font = regular_font
            c.number_format = money_format
            c.border = thin_border
        
        # Total Fila formula
        end_letter = get_column_letter(len(items) + 1)
        tot_cell = ws1.cell(row=r_offset, column=total_col_idx, value=f"=SUM(B{r_offset}:{end_letter}{r_offset})")
        tot_cell.font = bold_font
        tot_cell.number_format = money_format
        tot_cell.border = thin_border
    
    # Total Columnas
    tot_row_idx = 13
    ws1.cell(row=tot_row_idx, column=1, value="Total").font = bold_font
    for col_idx in range(2, len(items) + 2):
        col_letter = get_column_letter(col_idx)
        c = ws1.cell(row=tot_row_idx, column=col_idx, value=f"=SUM({col_letter}8:{col_letter}12)")
        c.font = bold_font
        c.number_format = money_format
        c.border = thin_border
        
    tot_general = ws1.cell(row=tot_row_idx, column=total_col_idx, value=f"=SUM({get_column_letter(total_col_idx)}8:{get_column_letter(total_col_idx)}12)")
    tot_general.font = bold_font
    tot_general.number_format = money_format
    tot_general.fill = PatternFill(start_color="FEF08A", end_color="FEF08A", fill_type="solid")
    tot_general.border = thin_border
    
    # Auto-adjust column widths
    for col in ws1.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws1.column_dimensions[col_letter].width = max(max_len + 3, 13)
        
    # ----------------------------------------------------
    # HOJA 2: Consumo por Integrante
    # ----------------------------------------------------
    ws2 = wb.create_sheet(title="Consumo por Integrante")
    
    integrantes = [
        "Pamela", "Cote", "Joaquín", "Ignacio", "Quena",
        "Ale", "Mauri", "Mindy", "Monse", "Gustavo",
        "Coni", "Miriam", "Abuelo Coni", "Mamá Mauri", "Papá Mauri",
        "Franklin"
    ]
    
    ws2["A1"] = "Integrante"
    ws2["A1"].font = header_font
    ws2["A1"].fill = header_fill
    ws2["A1"].alignment = Alignment(horizontal="center")
    
    for col_idx, item in enumerate(items, start=2):
        c = ws2.cell(row=1, column=col_idx, value=item)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center")
        
    si_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid") # green
    no_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid") # red
    
    for r_idx, persona in enumerate(integrantes, start=2):
        ws2.cell(row=r_idx, column=1, value=persona).font = bold_font
        for col_idx, item in enumerate(items, start=2):
            item_num = col_idx - 1 # 1 to 19
            # Items 1 to 13 son asado
            # Items 14 a 19 son empanadas:
            # 14: Pino (1a), 15: Pino (2a), 16: Mechada, 17: Camarón Queso, 18: Queso, 19: Napolitana
            if item_num <= 13:
                if persona == "Monse":
                    is_si = (item_num >= 10)
                elif persona == "Miriam":
                    is_si = (item_num not in [1, 2, 3, 9])
                else:
                    is_si = True
            else:
                # Empanadas
                if item == "Pino (1a Empanada)":
                    is_si = persona in ["Joaquín", "Ignacio", "Gustavo", "Mauri", "Pamela", "Quena", "Mindy"]
                elif item == "Pino (2a Empanada)":
                    is_si = persona in ["Joaquín", "Ignacio", "Gustavo", "Mauri", "Pamela", "Quena"]
                elif item == "Mechada":
                    is_si = persona in ["Ale", "Franklin"]
                elif item == "Camarón Queso":
                    is_si = persona == "Franklin"
                elif item == "Queso":
                    is_si = persona == "Cote"
                elif item == "Napolitana":
                    is_si = persona == "Cote"
                else:
                    is_si = False
                
            val_text = "Sí" if is_si else "No"
            c = ws2.cell(row=r_idx, column=col_idx, value=val_text)
            c.alignment = Alignment(horizontal="center")
            c.font = regular_font
            c.fill = si_fill if is_si else no_fill
            c.border = thin_border
            
    # Fila de conteo
    count_row = len(integrantes) + 2
    ws2.cell(row=count_row, column=1, value="Total Consumidores").font = bold_font
    for col_idx in range(2, len(items) + 2):
        col_letter = get_column_letter(col_idx)
        c = ws2.cell(row=count_row, column=col_idx, value=f'=COUNTIF({col_letter}2:{col_letter}{count_row-1}, "Sí")')
        c.font = bold_font
        c.alignment = Alignment(horizontal="center")
        c.border = thin_border
        
    for col in ws2.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws2.column_dimensions[col_letter].width = max(max_len + 3, 14)

    # ----------------------------------------------------
    # HOJA 3: División de Cuentas & Liquidación
    # ----------------------------------------------------
    ws3 = wb.create_sheet(title="División de Cuentas")
    
    headers_ws3 = [
        "Integrante", "Consumo Propio", "Aporte Cuota Quena (1/4)", 
        "Total Cuota a Pagar", "Aportado en Compras", "Saldo a Liquidar", 
        "Destinatario / Acción", "Estado"
    ]
    
    for c_idx, h in enumerate(headers_ws3, start=1):
        cell = ws3.cell(row=1, column=c_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        
    sponsors_quena = ["Pamela", "Ale", "Mindy", "Joaquín"]
    quena_row = integrantes.index("Quena") + 2
    
    for r_idx, persona in enumerate(integrantes, start=2):
        ws3.cell(row=r_idx, column=1, value=persona).font = bold_font
        
        # Col B: Consumo Propio
        formula_terms = []
        for col_idx in range(2, len(items) + 2):
            col_letter = get_column_letter(col_idx)
            term = f"IF('Consumo por Integrante'!{col_letter}{r_idx}=\"Sí\", Compras!{col_letter}$13 / 'Consumo por Integrante'!{col_letter}${count_row}, 0)"
            formula_terms.append(term)
        
        exact_formula = "=" + " + ".join(formula_terms)
        c_exact = ws3.cell(row=r_idx, column=2, value=f"=ROUND({exact_formula[1:]}, 0)")
        c_exact.font = regular_font
        c_exact.number_format = money_format
        c_exact.border = thin_border
        
        # Col C: Aporte Cuota Quena (1/4)
        if persona in sponsors_quena:
            c_quena = ws3.cell(row=r_idx, column=3, value=f"=ROUND(B${quena_row}/4, 0)")
            c_quena.font = Font(name="Calibri", size=11, color="1E3A8A", bold=True)
        elif persona == "Quena":
            c_quena = ws3.cell(row=r_idx, column=3, value=f"=-B{r_idx}")
            c_quena.font = Font(name="Calibri", size=11, color="DC2626", bold=True)
        else:
            c_quena = ws3.cell(row=r_idx, column=3, value=0)
            c_quena.font = regular_font
        c_quena.number_format = money_format
        c_quena.border = thin_border
        
        # Col D: Total Cuota a Pagar
        if persona == "Quena":
            c_tot_pay = ws3.cell(row=r_idx, column=4, value=0)
        else:
            c_tot_pay = ws3.cell(row=r_idx, column=4, value=f"=B{r_idx}+C{r_idx}")
        c_tot_pay.font = bold_font
        c_tot_pay.number_format = money_format
        c_tot_pay.border = thin_border
        
        # Col E: Aportado en Compras
        aporte = 0
        if persona == "Ignacio":
            aporte = 183392
        elif persona == "Mindy":
            aporte = 11790
        elif persona == "Joaquín":
            aporte = 51820
            
        c_paid = ws3.cell(row=r_idx, column=5, value=aporte)
        c_paid.font = regular_font
        c_paid.number_format = money_format
        c_paid.border = thin_border
        
        # Col F: Saldo a Liquidar
        # Col G: Destinatario / Acción
        # Col H: Estado
        if persona == "Ignacio":
            c_net = ws3.cell(row=r_idx, column=6, value=f"=E{r_idx}-B{r_idx}")
            c_net.font = Font(name="Calibri", size=11, bold=True, color="1E3A8A")
            c_dest = ws3.cell(row=r_idx, column=7, value="Organizador (Recauda $193.121)")
            c_status = ws3.cell(row=r_idx, column=8, value="Organizador")
        elif persona == "Quena":
            c_net = ws3.cell(row=r_idx, column=6, value=0)
            c_net.font = Font(name="Calibri", size=11, bold=True, color="6B7280")
            c_dest = ws3.cell(row=r_idx, column=7, value="Cubierto por familiares")
            c_status = ws3.cell(row=r_idx, column=8, value="Cubierto")
        elif persona == "Joaquín":
            # Saldo a favor
            c_net = ws3.cell(row=r_idx, column=6, value=f"=E{r_idx}-D{r_idx}")
            c_net.font = Font(name="Calibri", size=11, bold=True, color="2563EB")
            c_dest = ws3.cell(row=r_idx, column=7, value="Ignacio le transfiere (Saldo a favor)")
            c_status = ws3.cell(row=r_idx, column=8, value="A favor (Recibe)")
        else:
            c_net = ws3.cell(row=r_idx, column=6, value=f"=IF(D{r_idx}-E{r_idx}>0, D{r_idx}-E{r_idx}, 0)")
            c_net.font = Font(name="Calibri", size=11, bold=True, color="047857")
            c_dest = ws3.cell(row=r_idx, column=7, value="Transfiere a Ignacio")
            c_status = ws3.cell(row=r_idx, column=8, value="Pendiente")
            
        c_net.number_format = money_format
        c_net.border = thin_border
        
        c_dest.font = regular_font
        c_dest.alignment = Alignment(horizontal="center")
        c_dest.border = thin_border
        
        c_status.font = regular_font
        c_status.alignment = Alignment(horizontal="center")
        c_status.border = thin_border
        
    tot_row = len(integrantes) + 2
    ws3.cell(row=tot_row, column=1, value="Totales").font = bold_font
    
    t_propio = ws3.cell(row=tot_row, column=2, value=f"=SUM(B2:B{tot_row-1})")
    t_propio.font = bold_font
    t_propio.number_format = money_format
    t_propio.fill = PatternFill(start_color="FEF08A", end_color="FEF08A", fill_type="solid")
    t_propio.border = thin_border
    
    t_quena = ws3.cell(row=tot_row, column=3, value=f"=SUM(C2:C{tot_row-1})")
    t_quena.font = bold_font
    t_quena.number_format = money_format
    t_quena.fill = PatternFill(start_color="FEF08A", end_color="FEF08A", fill_type="solid")
    t_quena.border = thin_border

    t_tot_pay = ws3.cell(row=tot_row, column=4, value=f"=SUM(D2:D{tot_row-1})")
    t_tot_pay.font = bold_font
    t_tot_pay.number_format = money_format
    t_tot_pay.fill = PatternFill(start_color="FEF08A", end_color="FEF08A", fill_type="solid")
    t_tot_pay.border = thin_border

    t_paid = ws3.cell(row=tot_row, column=5, value=f"=SUM(E2:E{tot_row-1})")
    t_paid.font = bold_font
    t_paid.number_format = money_format
    t_paid.fill = PatternFill(start_color="FEF08A", end_color="FEF08A", fill_type="solid")
    t_paid.border = thin_border

    t_net = ws3.cell(row=tot_row, column=6, value=f"=SUMIF(G2:G{tot_row-1}, \"Transfiere a Ignacio\", F2:F{tot_row-1})")
    t_net.font = bold_font
    t_net.number_format = money_format
    t_net.fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    t_net.border = thin_border

    ws3.cell(row=tot_row, column=7, value="Total Recaudado por Ignacio")
    ws3.cell(row=tot_row, column=7).font = bold_font
    ws3.cell(row=tot_row, column=8, value="")

    for col in ws3.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws3.column_dimensions[col_letter].width = max(max_len + 4, 16)
        
    wb.save(file_path)
    print(f"Excel guardado con éxito en: {file_path}")

if __name__ == "__main__":
    update_excel("Cuentas Fiestas Patrias.xlsx")
    try:
        update_excel("C:/Users/Ignac/OneDrive/Escritorio/Cuentas Fiestas Patrias.xlsx")
    except Exception as e:
        print("No se pudo actualizar la copia en el escritorio:", e)
